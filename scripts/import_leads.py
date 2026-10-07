"""Load Amitek_Lead_Sheet.xlsx into the bot database so the bot knows who is writing.

Usage: DATABASE_URL=... python -m scripts.import_leads path/to/Amitek_Lead_Sheet.xlsx
"""
import sys

from openpyxl import load_workbook

from app.config import settings
from app.db import Store, make_engine

TIER_MAP = {"Applicator / Project": "Applicator / Project", "Dealer": "Dealer", "Bulk": "Bulk",
            "End Client": "End Client"}


def main(path: str) -> None:
    store = Store(make_engine(settings.database_url))
    ws = load_workbook(path, read_only=True)["Leads"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    n = skipped = 0
    for values in rows:
        r = dict(zip(header, values))
        phone = str(r.get("WhatsApp Number") or "").strip()
        if not phone.isdigit():
            skipped += 1
            continue
        existing = store.get_lead(phone)
        if existing and existing.get("status") not in (None, "New"):
            skipped += 1  # never overwrite a lead the bot is already talking to
            continue
        store.upsert_lead(
            phone,
            lead_id=r.get("Lead ID"),
            business_name=r.get("Display Name") or r.get("Business Name"),
            category=r.get("Lead Category") or "Other",
            tier=TIER_MAP.get(r.get("Customer Tier"), "To confirm"),
            city=r.get("City"), state=r.get("State"),
            campaign=r.get("Campaign") or None,
            opt_in=r.get("Opt-in Status") or "Not Asked",
            status=r.get("Lead Status") or "New",
        )
        n += 1
    print(f"imported {n}, skipped {skipped}")


if __name__ == "__main__":
    main(sys.argv[1])
